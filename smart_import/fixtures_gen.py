"""Genera archivos de prueba que imitan exports reales de clientes.

No se versionan binarios: se regeneran con `python -m smart_import make-fixtures`.
Cada fixture apunta a UN problema concreto del pipeline.
"""
from __future__ import annotations

import csv
import json
import random
from pathlib import Path

AR_STREETS = ["Av. Corrientes", "Av. Santa Fe", "Av. Cabildo", "Maipú", "Florida", "Lavalle",
              "Tucumán", "Sarmiento", "Esmeralda", "Reconquista", "San Martín", "Viamonte",
              "Cerrito", "Libertad", "Paraguay", "Av. del Libertador", "Callao", "Rivadavia"]
AR_ZONES = ["Retiro", "San Nicolás", "Recoleta", "Palermo", "Belgrano", "Almagro"]
US_STREETS = ["Broadway", "Pine St", "Madison St", "E Pike St", "Boylston Ave", "E Denny Way",
              "1st Ave", "Union St", "Olive Way", "Yesler Way"]
US_ZONES = ["C-1", "C-2", "B-1", "G-3", "N-2"]
FIRST = ["Juan", "Maria", "Carlos", "Lucia", "Martin", "Sofia", "Diego", "Ana", "Pablo", "Valeria"]
LAST = ["Pérez", "Gómez", "López", "Fernández", "Rodríguez", "Martínez", "Sosa", "Díaz"]
US_FIRST = ["James", "Mary", "Robert", "Linda", "Michael", "Sarah", "David", "Emily"]
US_LAST = ["Smith", "Johnson", "Brown", "Davis", "Miller", "Wilson", "Moore", "Taylor"]
PACKAGING = ["BOX", "BOX", "BOX", "BAG", "PALLET"]


def _rng(seed: int) -> random.Random:
    return random.Random(seed)


def _stop_ar(r: random.Random, i: int) -> dict:
    return {
        "delivery_id": f"VP-{2000 + i}",
        "lat": round(-34.60 + r.uniform(-0.05, 0.05), 6),
        "lng": round(-58.38 + r.uniform(-0.05, 0.05), 6),
        "address": f"{r.choice(AR_STREETS)} {r.randrange(100, 4000)}, Buenos Aires",
        "zone": r.choice(AR_ZONES),
        "customer": f"{r.choice(FIRST)} {r.choice(LAST)}",
        "phone": f"11{r.randrange(1000, 9999)}{r.randrange(1000, 9999)}",
        "packages": r.choices([1, 1, 1, 2, 3], k=1)[0],
        "weight": round(r.uniform(0.05, 4.5), 2),
        "length": round(r.uniform(15, 40), 1),
        "width": round(r.uniform(10, 30), 1),
        "height": round(r.uniform(4, 20), 1),
        "priority": r.randrange(1, 4),
        "service": r.choice([5, 10, 15, 20]),
    }


def _stop_us(r: random.Random, i: int) -> dict:
    return {
        "delivery_id": f"DLV-{2000 + i}",
        "lat": round(47.61 + r.uniform(-0.04, 0.04), 6),
        "lng": round(-122.29 + r.uniform(-0.06, 0.06), 6),
        "address": f"{r.randrange(100, 1500)} {r.choice(US_STREETS)}, Seattle",
        "zone": r.choice(US_ZONES),
        "customer": f"{r.choice(US_FIRST)} {r.choice(US_LAST)}",
        "phone": f"206{r.randrange(200, 999)}{r.randrange(1000, 9999)}",
        "packages": r.choices([1, 1, 2, 3], k=1)[0],
        "weight": round(r.uniform(0.1, 6.0), 2),
        "length": round(r.uniform(15, 40), 1),
        "width": round(r.uniform(10, 30), 1),
        "height": round(r.uniform(4, 20), 1),
        "priority": r.randrange(1, 4),
        "service": r.choice([5, 10, 15]),
    }


def _write_csv(path: Path, header: list[str], rows: list[list], delimiter: str = ",",
               encoding: str = "utf-8", preamble: list[list] | None = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding=encoding, newline="", errors="replace") as fh:
        w = csv.writer(fh, delimiter=delimiter, lineterminator="\n")
        for line in preamble or []:
            w.writerow(line)
        w.writerow(header)
        w.writerows(rows)
    return path


def _write_xlsx(path: Path, header: list[str], rows: list[list],
                sheet: str = "Hoja1", preamble: list[list] | None = None,
                extra_sheets: list[str] | None = None) -> Path:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for line in preamble or []:
        ws.append(line)
    ws.append(header)
    for row in rows:
        ws.append(row)
    for name in extra_sheets or []:
        wb.create_sheet(title=name).append(["hoja de notas, no son datos"])
    wb.save(path)
    return path


# ---------------------------------------------------------------- fixtures

def es_headers_raros(out: Path, n: int, seed: int = 1) -> Path:
    """Headers abreviados en espanol. El caso tipico del cliente argentino."""
    r = _rng(seed)
    header = ["Dest.", "Receptor", "Contacto 1", "Kg", "Cant bultos", "Latitud", "Longitud",
              "Zona", "Obs"]
    rows = []
    for i in range(n):
        s = _stop_ar(r, i)
        rows.append([s["address"], s["customer"], s["phone"], s["weight"], s["packages"],
                     s["lat"], s["lng"], s["zone"], r.choice(["", "Portero electrico", "Local rojo"])])
    return _write_xlsx(out / "es_headers_raros.xlsx", header, rows, sheet="Entregas")


def es_sin_coords(out: Path, n: int, seed: int = 2) -> Path:
    """Sin lat/lng: TODAS las filas quedan en needs_geocode."""
    r = _rng(seed)
    header = ["Domicilio Entrega", "Cliente", "Cel", "Bultos", "Peso kg", "Referencia"]
    rows = []
    for i in range(n):
        s = _stop_ar(r, i)
        rows.append([s["address"], s["customer"], s["phone"], s["packages"], s["weight"],
                     r.choice(["", "Timbre 3B", "Dejar en porteria"])])
    return _write_csv(out / "es_sin_coords.csv", header, rows)


def en_weird(out: Path, n: int, seed: int = 3) -> Path:
    r = _rng(seed)
    header = ["Ship To Address", "Recipient", "Mobile", "Pkg Count", "Weight (kg)",
              "Latitude", "Longitude", "Delivery Notes"]
    rows = []
    for i in range(n):
        s = _stop_us(r, i)
        rows.append([s["address"], s["customer"], s["phone"], s["packages"], s["weight"],
                     s["lat"], s["lng"], r.choice(["", "Leave at door", "Ring bell"])])
    return _write_csv(out / "en_weird.csv", header, rows)


def semicolon_latin1(out: Path, n: int, seed: int = 4) -> Path:
    """Export tipico de Excel es-AR: separador ';' , coma decimal y cp1252."""
    r = _rng(seed)
    header = ["Dirección", "Cliente", "Teléfono", "Peso", "Latitud", "Longitud"]
    rows = []
    for i in range(n):
        s = _stop_ar(r, i)
        rows.append([s["address"], s["customer"], s["phone"],
                     f"{s['weight']:.2f}".replace(".", ","),
                     f"{s['lat']:.6f}".replace(".", ","),
                     f"{s['lng']:.6f}".replace(".", ",")])
    return _write_csv(out / "semicolon_latin1.csv", header, rows,
                      delimiter=";", encoding="cp1252")


def pipe_delimited(out: Path, n: int, seed: int = 5) -> Path:
    r = _rng(seed)
    header = ["direccion", "cliente", "telefono", "lat", "lon", "bultos"]
    rows = [[(s := _stop_ar(r, i))["address"], s["customer"], s["phone"],
             s["lat"], s["lng"], s["packages"]] for i in range(n)]
    return _write_csv(out / "pipe_delimited.txt", header, rows, delimiter="|")


def tabs_tsv(out: Path, n: int, seed: int = 6) -> Path:
    r = _rng(seed)
    header = ["address", "customer_name", "phone", "lat", "lng", "weight_kg"]
    rows = [[(s := _stop_us(r, i))["address"], s["customer"], s["phone"],
             s["lat"], s["lng"], s["weight"]] for i in range(n)]
    return _write_csv(out / "tabs.tsv", header, rows, delimiter="\t")


def preamble_dirty(out: Path, n: int, seed: int = 7) -> Path:
    """Titulo, logo y fecha antes del header + filas vacias + columnas basura."""
    r = _rng(seed)
    preamble = [["REPORTE DE ENTREGAS - LOGISTICA ACME S.A."], [],
                ["Generado:", "2026-09-15", "", "Usuario:", "operaciones"], []]
    header = ["Dir. entrega", "Nom Dest", "Tel dest", "Latitud", "Longitud",
              "Cant.", "Kg", "Codigo interno", "Sucursal origen"]
    rows: list[list] = []
    for i in range(n):
        s = _stop_ar(r, i)
        rows.append([s["address"], s["customer"], s["phone"], s["lat"], s["lng"],
                     s["packages"], s["weight"], f"INT{r.randrange(10000, 99999)}", "CD-NORTE"])
        if i and i % 7 == 0:
            rows.append([None] * len(header))              # fila vacia intercalada
    return _write_xlsx(out / "preamble_dirty.xlsx", header, rows, sheet="Datos",
                       preamble=preamble, extra_sheets=["Instructivo", "Resumen"])


def swapped_coords(out: Path, n: int, seed: int = 8) -> Path:
    """lat y lng invertidas: el sistema debe AVISAR, no corregir en silencio."""
    r = _rng(seed)
    header = ["address", "lat", "lng", "customer_name"]
    rows = []
    for i in range(n):
        s = _stop_us(r, i)
        rows.append([s["address"], s["lng"], s["lat"], s["customer"]])   # <- invertidas
    return _write_csv(out / "swapped_coords.csv", header, rows)


def one_row_per_delivery(out: Path, n: int, seed: int = 9) -> Path:
    """'bultos: 3' sin package_id -> hay que expandir a 3 bultos."""
    r = _rng(seed)
    header = ["delivery_id", "address", "lat", "lng", "cantidad de bultos", "peso kg", "cliente"]
    rows = []
    for i in range(n):
        s = _stop_ar(r, i)
        rows.append([s["delivery_id"], s["address"], s["lat"], s["lng"],
                     s["packages"], s["weight"], s["customer"]])
    return _write_xlsx(out / "one_row_per_delivery.xlsx", header, rows, sheet="deliveries")


def merged_field(out: Path, n: int, seed: int = 10) -> Path:
    """Todo en una sola columna de texto libre: el caso que necesita el modelo."""
    r = _rng(seed)
    header = ["Datos entrega", "Zona"]
    rows = []
    for i in range(n):
        s = _stop_ar(r, i)
        blob = r.choice([
            f"{s['customer']} - {s['address']} - tel {s['phone']}",
            f"{s['address']} ({s['customer']}, cel {s['phone']})",
            f"Entregar a {s['customer']} en {s['address']}. Telefono: {s['phone']}",
        ])
        rows.append([blob, s["zone"]])
    return _write_csv(out / "merged_field.csv", header, rows)


def no_headers(out: Path, n: int, seed: int = 11) -> Path:
    """Headers inutiles (Campo 1..N): solo las heuristicas de contenido salvan esto."""
    r = _rng(seed)
    header = [f"Campo {i}" for i in range(1, 7)]
    rows = []
    for i in range(n):
        s = _stop_ar(r, i)
        rows.append([s["delivery_id"], s["address"], s["lat"], s["lng"], s["phone"], s["customer"]])
    return _write_csv(out / "no_headers.csv", header, rows)


def mixed_locale(out: Path, n: int, seed: int = 12) -> Path:
    r = _rng(seed)
    header = ["address", "lat", "lng", "customer_name", "phone", "weight_kg"]
    rows = []
    for i in range(n):
        s = _stop_ar(r, i) if i % 2 == 0 else _stop_us(r, i)
        rows.append([s["address"], s["lat"], s["lng"], s["customer"], s["phone"], s["weight"]])
    return _write_xlsx(out / "mixed_locale.xlsx", header, rows, sheet="deliveries")


def legacy_xls(out: Path, n: int, seed: int = 13) -> Path | None:
    """XLS binario viejo (BIFF). Se omite si xlwt no esta instalado."""
    try:
        import xlwt
    except ImportError:
        return None
    r = _rng(seed)
    header = ["Direccion", "Cliente", "Latitud", "Longitud", "Bultos"]
    wb = xlwt.Workbook()
    ws = wb.add_sheet("deliveries")
    for c, name in enumerate(header):
        ws.write(0, c, name)
    for i in range(n):
        s = _stop_ar(r, i)
        for c, v in enumerate([s["address"], s["customer"], s["lat"], s["lng"], s["packages"]]):
            ws.write(i + 1, c, v)
    path = out / "legacy.xls"
    wb.save(str(path))
    return path


def scale(out: Path, n: int, seed: int = 99) -> Path:
    """Volumen, con las columnas raras del fixture es_headers_raros."""
    r = _rng(seed)
    header = ["Dest.", "Receptor", "Contacto 1", "Kg", "Cant bultos", "Latitud", "Longitud"]
    rows = []
    for i in range(n):
        s = _stop_ar(r, i)
        rows.append([s["address"], s["customer"], s["phone"], s["weight"], s["packages"],
                     s["lat"], s["lng"]])
    label = f"{n // 1000}k" if n >= 1000 else str(n)
    return _write_xlsx(out / f"scale_{label}.xlsx", header, rows, sheet="deliveries")


BUILDERS = {
    "es_headers_raros": es_headers_raros,
    "es_sin_coords": es_sin_coords,
    "en_weird": en_weird,
    "semicolon_latin1": semicolon_latin1,
    "pipe_delimited": pipe_delimited,
    "tabs": tabs_tsv,
    "preamble_dirty": preamble_dirty,
    "swapped_coords": swapped_coords,
    "one_row_per_delivery": one_row_per_delivery,
    "merged_field": merged_field,
    "no_headers": no_headers,
    "mixed_locale": mixed_locale,
    "legacy_xls": legacy_xls,
}


def build_all(out_dir: str | Path, rows: int = 40, sizes: tuple[int, ...] = ()) -> list[Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    created = [p for fn in BUILDERS.values() if (p := fn(out, rows)) is not None]
    created += [scale(out, n) for n in sizes]
    return created
