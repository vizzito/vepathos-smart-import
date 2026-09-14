"""XLSX (openpyxl, read_only) y XLS (xlrd)."""
from __future__ import annotations

from pathlib import Path
from typing import Any
from zipfile import ZipFile

MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 10000
MAX_COLUMNS = 1024
MAX_SCANNED_ROWS = 1000000


def _check_archive(path):
    # Shared strings/styles are loaded even with openpyxl read_only=True.
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_MEMBERS or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
            raise ValueError("Excel excede el límite de expansión permitido")

from .base import FileMeta, Table, choose_header_row, dedupe_columns, is_blank


def _clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        v = value.strip()
        return v or None
    return value


def _pick_sheet(names: list[str], wanted: str | None) -> str:
    if not names:
        raise ValueError("Excel sin hojas de datos")
    if wanted and wanted in names:
        return wanted
    for preferred in ("deliveries", "entregas", "stops", "data"):
        for n in names:
            if n.strip().lower() == preferred:
                return n
    return names[0]


def read_xlsx(path: str | Path, max_rows: int | None = None, sheet: str | None = None) -> Table:
    from openpyxl import load_workbook

    p = Path(path)
    _check_archive(p)
    # read_only=True: openpyxl no materializa la hoja entera en memoria.
    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        sheets = list(wb.sheetnames)
        chosen = _pick_sheet(sheets, sheet)
        ws = wb[chosen]

        if (ws.max_column or 0) > MAX_COLUMNS:
            raise ValueError("Excel excede el límite de columnas permitido")
        head: list[list[Any]] = []
        it = ws.iter_rows(values_only=True, max_col=min(ws.max_column or MAX_COLUMNS, MAX_COLUMNS),
                          max_row=min(ws.max_row or MAX_SCANNED_ROWS, MAX_SCANNED_ROWS))
        for raw in it:
            head.append([_clean(c) for c in raw])
            if len(head) >= 12:
                break

        header_row = choose_header_row(head)
        columns = dedupe_columns(head[header_row] if head else [])
        width = len(columns)

        rows: list[tuple] = []

        def push(raw_row: list[Any]) -> bool:
            if all(is_blank(c) for c in raw_row):
                return True
            vals = list(raw_row[:width]) + [None] * max(0, width - len(raw_row))
            rows.append(tuple(vals))
            return not (max_rows and len(rows) >= max_rows)

        for raw_row in head[header_row + 1:]:
            if not push(raw_row):
                break
        else:
            for raw in it:
                if not push([_clean(c) for c in raw]):
                    break

        meta = FileMeta(
            path=str(p), format="xlsx", size_bytes=p.stat().st_size,
            sheet=chosen, sheets=sheets,
            header_row=header_row, preamble_rows=header_row,
        )
        if len(sheets) > 1:
            meta.notes.append(f"{len(sheets)} hojas; se leyo '{chosen}'")
        if header_row:
            meta.notes.append(f"{header_row} fila(s) de preambulo descartadas antes del header")
        return Table(meta=meta, columns=columns, rows=rows)
    finally:
        wb.close()


def read_xls(path: str | Path, max_rows: int | None = None, sheet: str | None = None) -> Table:
    import xlrd

    p = Path(path)
    book = xlrd.open_workbook(str(p), on_demand=True)
    try:
        sheets = book.sheet_names()
        chosen = _pick_sheet(sheets, sheet)
        sh = book.sheet_by_name(chosen)
        if sh.ncols > MAX_COLUMNS:
            raise ValueError("Excel excede el límite de columnas permitido")
        head = [[_clean(sh.cell_value(r, c)) for c in range(sh.ncols)]
                for r in range(min(12, sh.nrows))]
        header_row = choose_header_row(head)
        columns = dedupe_columns(head[header_row] if head else [])
        rows: list[tuple] = []
        for r in range(header_row + 1, min(sh.nrows, MAX_SCANNED_ROWS)):
            row = tuple(_clean(sh.cell_value(r, c)) for c in range(len(columns)))
            if all(is_blank(c) for c in row):
                continue
            rows.append(row)
            if max_rows and len(rows) >= max_rows:
                break
        meta = FileMeta(
            path=str(p), format="xls", size_bytes=p.stat().st_size,
            sheet=chosen, sheets=sheets, header_row=header_row, preamble_rows=header_row,
        )
        return Table(meta=meta, columns=columns, rows=rows)
    finally:
        book.release_resources()
